"""
Cross-validation. The charter's rule: if SQL, Excel and Python disagree on a
headline number, the analysis stops until the disagreement is explained.

Reads what Excel actually computed, recomputes the same measures in SQL and in
pandas over the identical scope, compares, and writes the comparison back into
the workbook's Validation sheet.
"""

import pandas as pd
import psycopg2
import warnings
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

warnings.filterwarnings("ignore")
WB = "excel/permit_review_analysis.xlsx"
TOL = 0.005          # 0.5% relative tolerance

SCOPE = """
FROM permits.fact_permits f
JOIN permits.dim_date da       ON da.date_key       = f.application_date_key
JOIN permits.dim_permit_type t ON t.permit_type_key = f.permit_type_key
JOIN permits.dim_geography g   ON g.geography_key   = f.geography_key
WHERE f.is_review_analysable
  AND f.days_to_admin_close BETWEEN 0 AND 365
  AND da.year BETWEEN 2017 AND 2025
  AND t.permit_family = 'Building'
"""

# ------------------------------------------------------------------ EXCEL
print("reading Excel computed values...")
wbv = load_workbook(WB, data_only=True)
k = wbv["KPI Summary"]
excel = {
    "Records in scope":        k["B5"].value,
    "Distinct permits":        k["B6"].value,
    "Mean days to issue":      k["B7"].value,
    "Approved within 30 days": k["B8"].value,
    "Took over 90 days":       k["B9"].value,
    "Longest review (days)":   k["B10"].value,
    "Records reporting cost":  k["B11"].value,
    "Mean reported cost":      k["B12"].value,
}
wbv.close()

# -------------------------------------------------------------------- SQL
print("recomputing in SQL...")
con = psycopg2.connect(dbname="toronto_permits", user="root",
                       host="/var/run/postgresql")
sql = pd.read_sql(f"""
SELECT COUNT(*)                                                      AS records,
       COUNT(DISTINCT f.permit_num)                                  AS permits,
       AVG(f.days_to_issue)                                          AS mean_days,
       AVG(CASE WHEN f.days_to_issue <= 30 THEN 1.0 ELSE 0 END)      AS within30,
       AVG(CASE WHEN f.days_to_issue >  90 THEN 1.0 ELSE 0 END)      AS over90,
       MAX(f.days_to_issue)                                          AS max_days,
       COUNT(*) FILTER (WHERE f.cost_is_reported)                    AS cost_n,
       AVG(f.est_const_cost) FILTER (WHERE f.cost_is_reported)       AS mean_cost
{SCOPE}""", con).iloc[0]

# ----------------------------------------------------------------- PYTHON
print("recomputing in pandas...")
df = pd.read_sql(f"""
SELECT f.permit_num, f.days_to_issue, f.est_const_cost, f.cost_is_reported
{SCOPE}""", con)
con.close()

py = {
    "records": len(df),
    "permits": df.permit_num.nunique(),
    "mean_days": df.days_to_issue.mean(),
    "within30": (df.days_to_issue <= 30).mean(),
    "over90": (df.days_to_issue > 90).mean(),
    "max_days": df.days_to_issue.max(),
    "cost_n": int(df.cost_is_reported.sum()),
    "mean_cost": df.loc[df.cost_is_reported, "est_const_cost"].mean(),
}

# ------------------------------------------------------------------ COMPARE
pairs = [
    ("Records in scope",        "records",   '#,##0'),
    ("Distinct permits",        "permits",   '#,##0'),
    ("Mean days to issue",      "mean_days", '#,##0.00'),
    ("Approved within 30 days", "within30",  '0.00%'),
    ("Took over 90 days",       "over90",    '0.00%'),
    ("Longest review (days)",   "max_days",  '#,##0'),
    ("Records reporting cost",  "cost_n",    '#,##0'),
    ("Mean reported cost",      "mean_cost", '$#,##0.00'),
]

results, all_ok = [], True
print("\n" + "=" * 92)
print(f"{'MEASURE':<26}{'EXCEL':>17}{'SQL':>17}{'PYTHON':>17}{'':>8}")
print("=" * 92)
for label, key, fmt in pairs:
    e = float(excel[label])
    s = float(sql[key])
    p = float(py[key])
    scale = max(abs(e), abs(s), abs(p), 1e-9)
    ok = (abs(e - s) / scale < TOL) and (abs(e - p) / scale < TOL)
    all_ok &= ok
    print(f"{label:<26}{e:>17,.4f}{s:>17,.4f}{p:>17,.4f}"
          f"{'   MATCH' if ok else '   ***MISMATCH***':>8}")
    results.append((label, e, s, p, "MATCH" if ok else "MISMATCH", fmt))
print("=" * 92)
print(("ALL MEASURES RECONCILE" if all_ok
       else "RECONCILIATION FAILED - analysis must stop") + "\n")

# ------------------------------------------------- write back to workbook
FONT, THIN = "Arial", Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
OKF = PatternFill("solid", fgColor="C6EFCE")
BADF = PatternFill("solid", fgColor="FFC7CE")

wb = load_workbook(WB)
ws = wb["Validation"]
for row in ws["A5:E40"]:
    for c in row:
        c.value, c.fill, c.border = None, PatternFill(), Border()

for i, (label, e, s, p, verdict, fmt) in enumerate(results, start=5):
    ws.cell(row=i, column=1, value=label).font = Font(name=FONT, size=10)
    for col, v in [(2, e), (3, s), (4, p)]:
        c = ws.cell(row=i, column=col, value=v)
        c.font, c.number_format, c.border = Font(name=FONT, size=10), fmt, BOX
    c = ws.cell(row=i, column=5, value=verdict)
    c.font = Font(name=FONT, bold=True, size=10,
                  color="006100" if verdict == "MATCH" else "9C0006")
    c.fill = OKF if verdict == "MATCH" else BADF
    c.border, c.alignment = BOX, Alignment(horizontal="center")

r = 5 + len(results) + 1
ws.cell(row=r, column=1, value="OVERALL").font = Font(name=FONT, bold=True, size=11)
c = ws.cell(row=r, column=5,
            value="ALL RECONCILE" if all_ok else "FAILED")
c.font = Font(name=FONT, bold=True, size=11,
              color="006100" if all_ok else "9C0006")
c.fill = OKF if all_ok else BADF
c.alignment = Alignment(horizontal="center")

for j, note in enumerate([
    "",
    "Excel column: values the workbook's own formulas produced (read with data_only).",
    "SQL column: recomputed by aggregate query against the PostgreSQL star schema.",
    "Python column: recomputed with pandas over the same extract.",
    f"Tolerance: {TOL:.1%} relative. Percentages are stored as fractions.",
    "Three independent paths from the same source. Agreement is evidence the logic is right,",
    "not proof - all three could share a wrong assumption. It does rule out transcription",
    "and single-tool errors, which is what it is for.",
], start=r + 2):
    ws.cell(row=j, column=1, value=note).font = Font(name=FONT, italic=True,
                                                     size=9, color="595959")
wb.save(WB)
print(f"validation written to {WB}")
